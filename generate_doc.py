import docx

doc = docx.Document()
doc.add_heading('ShadowPulse Dashboard Graph Documentation', 0)

doc.add_paragraph('This document outlines the various graphs used in the ShadowPulse dashboard, what data they represent, and their analytical use cases for system monitoring and threat detection.')

# 1. Telemetry Dashboard
doc.add_heading('1. Telemetry Dashboard', level=1)
doc.add_paragraph('The Telemetry dashboard provides an overview of the system\'s real-time hardware and OS metrics.')

doc.add_heading('CPU & Memory Utilization (Line Chart)', level=2)
doc.add_paragraph('Data Shown: Real-time CPU usage percentage and Memory usage percentage plotted against time.')
doc.add_paragraph('Use Case: Helps identify performance bottlenecks, resource exhaustion, or unusual spikes that may indicate unauthorized mining activity or a runaway process.')

doc.add_heading('Disk Activity (Line/Area Chart)', level=2)
doc.add_paragraph('Data Shown: Read and Write speeds (MB/s) over time.')
doc.add_paragraph('Use Case: Detects large data exfiltration, ransomware encryption activity, or heavy I/O operations by highlighting abnormal spikes in disk read/write rates.')

# 2. Network Dashboard
doc.add_heading('2. Network Dashboard', level=1)
doc.add_paragraph('The Network dashboard focuses on active connections and bandwidth usage.')

doc.add_heading('Network Traffic (Line Chart)', level=2)
doc.add_paragraph('Data Shown: Network inbound (Rx) and outbound (Tx) traffic (MB/s) over time.')
doc.add_paragraph('Use Case: Useful for spotting sudden data exfiltration (high Tx spike), incoming DDoS attempts, or suspicious communication with external C2 servers.')

doc.add_heading('Active Connections (Gauge or Stat Card)', level=2)
doc.add_paragraph('Data Shown: Total number of active network connections.')
doc.add_paragraph('Use Case: Highlights potential port scanning, reverse shells, or botnet activity by showing an unexpected surge in the number of concurrent connections.')

# 3. ML Engine Dashboard
doc.add_heading('3. ML Engine Dashboard', level=1)
doc.add_paragraph('The ML Engine dashboard visualizes the anomaly detection model\'s behavior and decisions.')

doc.add_heading('Real-time Threat Score (Line Chart)', level=2)
doc.add_paragraph('Data Shown: The calculated Threat Score (0-100) plotted over time, alongside an Anomaly Threshold line.')
doc.add_paragraph('Use Case: Provides a visual timeline of system security. If the threat score crosses the threshold, it triggers an alert. Helps analysts see how quickly a threat escalated and when the system returned to normal.')

doc.add_heading('Anomaly Detection Distribution (Scatter/Bar Chart)', level=2)
doc.add_paragraph('Data Shown: Clusters of normal vs. anomalous data points across different features (e.g., CPU vs Network).')
doc.add_paragraph('Use Case: Helps explain WHY the ML model flagged an event as anomalous. It visualizes which specific metrics deviated significantly from the established baseline.')

# 4. Global Components
doc.add_heading('4. Alert Center', level=1)
doc.add_paragraph('While not exclusively a graph, the alert feed is a critical visual component.')

doc.add_heading('Recent Alerts (Timeline/List)', level=2)
doc.add_paragraph('Data Shown: A chronological list of recent alerts, including severity, timestamp, and the specific metrics that triggered the alert.')
doc.add_paragraph('Use Case: Allows security analysts to review historical security events, investigate past anomalies, and track the frequency of alerts.')

doc.save('ShadowPulse_Graphs_Documentation.docx')
print('Document saved as ShadowPulse_Graphs_Documentation.docx')
